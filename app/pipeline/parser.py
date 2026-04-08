"""Document parser — PDF (pdfplumber), DOCX (python-docx), TXT."""
import os
from app.observability.logger import get_logger

logger = get_logger("parser")

def parse_document(file_path: str) -> dict:
    ext = os.path.splitext(file_path)[1].lower()
    parsers = {".pdf": _parse_pdf, ".docx": _parse_docx, ".txt": _parse_txt}
    parser = parsers.get(ext)
    if parser is None:
        return {"text": "", "page_count": 0, "status": "error", "error": f"Unsupported file format: {ext}"}
    try:
        result = parser(file_path)
        if result["status"] == "success":
            logger.info("parse_success", extra={"extra_data": {
                "file_type": ext, "text_length": len(result["text"]),
                "page_count": result["page_count"],
            }})
        return result
    except Exception as e:
        logger.error(f"Parse failed for {file_path}: {e}")
        return {"text": "", "page_count": 0, "status": "error", "error": str(e)}

def _parse_pdf(file_path: str) -> dict:
    import pdfplumber
    all_text_parts = []
    with pdfplumber.open(file_path) as pdf:
        page_count = len(pdf.pages)
        for page in pdf.pages:
            page_text_parts = []
            tables = page.extract_tables()
            if tables:
                for table in tables:
                    table_text = _format_table(table)
                    if table_text.strip():
                        page_text_parts.append(table_text)
            text = page.extract_text() or ""
            if text.strip():
                page_text_parts.append(text)
            all_text_parts.append("\n".join(page_text_parts))
    full_text = "\n\n--- Page Break ---\n\n".join(all_text_parts)
    return {"text": full_text, "page_count": page_count, "status": "success"}

def _format_table(table: list[list]) -> str:
    if not table:
        return ""
    rows = []
    for row in table:
        cells = [str(cell).strip() if cell else "" for cell in row]
        rows.append(" | ".join(cells))
    return "\n".join(rows)

def _parse_docx(file_path: str) -> dict:
    from docx import Document
    doc = Document(file_path)
    paragraphs = []
    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text)
    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            table_rows.append(" | ".join(cells))
        paragraphs.append("\n".join(table_rows))
    return {"text": "\n\n".join(paragraphs), "page_count": 1, "status": "success"}

def _parse_txt(file_path: str) -> dict:
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()
    return {"text": text, "page_count": 1, "status": "success"}
